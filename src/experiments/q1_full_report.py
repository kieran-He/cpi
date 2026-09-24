"""Validate Q1 full-run checkpoints and publish tables/template without solving.

Run after q1_full.py: ``uv run --locked python -m src.experiments.q1_full_report``.
Only results/checkpoints/q1_full/*.json are authoritative inputs. The P1 slice
has its own directory and can never become a formal 80-box result here.
Missing, corrupted, paused, timed-out or unproven scenarios stay explicitly
incomplete in the summary. An XLSX is issued only for a proven baseline.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
from collections import Counter
from copy import copy
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from src.experiments.q1_full import (CHECKPOINT_DIR, DEM, ROOT, TABLE_DIR, atomic_json,
                                     file_hash, make_fingerprint, read_checkpoint, write_csv_atomic)
from src.models.q1_full import (OBJECTIVES, Scenario, all_scenarios, flight_energy,
                                load_full_inputs, make_geometries, max_safe_payload)

TEMPLATE = ROOT / "questions" / "Results Submission Template.xlsx"
XLSX_OUT = ROOT / "results" / "Q1_Results_Submission.xlsx"
STATUS_OUT = ROOT / "results" / "logs" / "q1_full_report_status.json"
Q1_SHEET = "Q1_Single-Destination Batching"
Q1_HEADERS = ("Sortie ID", "Service Area ID", "Model ID", "Cargo Box ID List",
              "Total Mass (kg)", "Total Volume (m³)", "Round-Trip Time (s)",
              "Sortie Energy (kWh)", "Return SOC (%)")

CSV_COLUMNS = {
    "q1_full_input_boxes.csv": ["Cargo Box ID", "Service Area ID", "Mass (kg)", "Volume (m³)"],
    "q1_full_input_areas.csv": ["Service Area ID", "Longitude (deg)", "Latitude (deg)", "Elevation (m)", "Cargo count"],
    "q1_full_capacity_all.csv": ["Scenario", "rho", "alpha", "beta", "Service Area ID", "Model ID", "Status", "Maximum Safe Payload (kg)"],
    "q1_full_scenario_status.csv": ["Scenario", "rho", "alpha", "beta", "Checkpoint status", "Report status", "Formal", "Reason", "Candidates possible", "Candidates physical", "Candidates retained", "Dominated removed", "Epsilon planned", "Epsilon optimal", "Epsilon infeasible", "Epsilon tentative"],
    "q1_full_endpoints.csv": ["Scenario", "Primary", "Status", "Sorties", "Energy (kWh)", "Operation Time (s)", "Primary lower bound", "Area count"],
    "q1_full_pareto_sample.csv": ["Scenario", "Sample ID", "Sample status", "Sorties", "Energy (kWh)", "Operation Time (s)", "Selected batch IDs"],
    "q1_full_epsilon_status.csv": ["Scenario", "Job ID", "Stage", "Primary", "Bound N", "Bound E (kWh)", "Bound T (s)", "Status", "Solver status", "Reused certificate", "Proof source", "Root proof source", "Sorties", "Energy (kWh)", "Operation Time (s)", "Primary lower bound", "MIP gap", "Nodes", "Selected batches"],
    "q1_full_sensitivity.csv": ["Scenario", "Report status", "Representative type", "Sorties", "Energy (kWh)", "Operation Time (s)", "Delta sorties", "Delta energy (kWh)", "Delta operation time (s)", "Changed batch pairs", "Baseline comparison available"],
    "q1_full_formal_batches.csv": ["Scenario", *Q1_HEADERS, "Operation Time (s)", "Batch ID"],
}


def _near(a: object, b: object, *, atol: float = 1e-7, rtol: float = 1e-7) -> bool:
    try:
        x, y = float(a), float(b)
    except (TypeError, ValueError):
        return False
    return math.isfinite(x) and math.isfinite(y) and math.isclose(x, y, abs_tol=atol, rel_tol=rtol)


def _solution(ids: list[str], cargo_groups, aircrafts, geometries, scenario: Scenario) -> tuple[dict, list[dict], set[tuple]]:
    """Rebuild the selected sorties from raw cargo IDs and cached DEM geometry."""
    if not isinstance(ids, list) or len(ids) != len(set(ids)):
        raise ValueError("Duplicate/invalid selected batch IDs")
    coverage = Counter()
    rows = []
    pairs = set()
    for batch_id in ids:
        parts = batch_id.split("|")
        if len(parts) != 3 or parts[0] not in cargo_groups or parts[1] not in aircrafts:
            raise ValueError(f"Invalid batch ID {batch_id}")
        area, model, mask_text = parts
        boxes = cargo_groups[area]
        try:
            mask = int(mask_text)
        except ValueError as exc:
            raise ValueError(f"Invalid subset mask {batch_id}") from exc
        if mask < 1 or mask >= (1 << len(boxes)) or batch_id != f"{area}|{model}|{mask:05d}":
            raise ValueError(f"Noncanonical subset mask {batch_id}")
        chosen = [boxes[i] for i in range(len(boxes)) if mask & (1 << i)]
        mass = math.fsum(x.mass_kg for x in chosen)
        volume = math.fsum(x.volume_m3 for x in chosen)
        a = aircrafts[model]
        if mass > a.payload_capacity_kg + 1e-8 or volume > a.loading_volume_m3 + 1e-8:
            raise ValueError(f"Capacity violation {batch_id}")
        energy, flight, range_fraction = flight_energy(a, geometries[area], mass, scenario)
        soc = 1. - energy / a.usable_energy_kwh
        if range_fraction > 1. + 1e-8 or soc < scenario.rho - 1e-8:
            raise ValueError(f"Range/SOC violation {batch_id}")
        hand = a.base_handover_time_s + len(chosen) * a.handover_time_per_box_s
        round_trip = flight + hand
        operation = round_trip + a.prep_time_s + len(chosen) * a.loading_time_per_box_s
        cargo_ids = tuple(x.cargo_id for x in chosen)
        coverage.update(cargo_ids)
        pairs.add((area, model, cargo_ids))
        rows.append({"Scenario": scenario.key, "Sortie ID": "", "Service Area ID": area,
                     "Model ID": model, "Cargo Box ID List": ", ".join(cargo_ids),
                     "Total Mass (kg)": mass, "Total Volume (m³)": volume,
                     "Round-Trip Time (s)": round_trip, "Sortie Energy (kWh)": energy,
                     "Return SOC (%)": 100. * soc, "Operation Time (s)": operation, "Batch ID": batch_id})
    expected = {c.cargo_id for items in cargo_groups.values() for c in items}
    if set(coverage) != expected or any(coverage[x] != 1 for x in expected):
        raise ValueError("80-box coverage is not exactly once")
    rows.sort(key=lambda row: row["Batch ID"])
    for index, row in enumerate(rows, 1):
        row["Sortie ID"] = f"Q1-{index:03d}"
    return {"N": len(rows), "E": math.fsum(r["Sortie Energy (kWh)"] for r in rows),
            "T": math.fsum(r["Operation Time (s)"] for r in rows)}, rows, pairs


def _assert_metrics(actual: dict, recorded: dict) -> None:
    if actual["N"] != recorded.get("N") or not _near(actual["E"], recorded.get("E")) or not _near(actual["T"], recorded.get("T")):
        raise ValueError(f"Metrics disagree: recomputed={actual}, checkpoint={recorded}")


def _expected_representative(points: list[dict], endpoints: dict) -> tuple[list[str], dict, dict]:
    """Recheck the Q1 equal-weight ideal-distance selection rule."""
    if not points:
        raise ValueError("No proven Pareto sample")
    ideals = {j: endpoints[j]["metrics"][j] for j in OBJECTIVES}
    upper = {j: max(p["metrics"][j] for p in points) for j in OBJECTIVES}
    def key(point: dict) -> tuple:
        metric = point["metrics"]
        distance = math.sqrt(sum(((metric[j] - ideals[j]) / (upper[j] - ideals[j])) ** 2
                                 for j in OBJECTIVES if upper[j] > ideals[j] + 1e-8))
        return (round(distance, 9), metric["N"], metric["E"], metric["T"], tuple(sorted(point["selected"])))
    chosen = min(points, key=key)
    return chosen["selected"], ideals, upper


def _capacity_rows(state: dict, scenario: Scenario, areas, aircrafts, geometries) -> list[dict]:
    entries = state.get("capacity")
    if not isinstance(entries, list) or len(entries) != 45:
        raise ValueError("Capacity table must contain 45 model-area rows")
    found = {}
    for entry in entries:
        area, model = entry.get("Service Area ID"), entry.get("Model ID")
        if area not in areas or model not in aircrafts or (area, model) in found:
            raise ValueError("Invalid/duplicate capacity model-area row")
        status, payload = max_safe_payload(aircrafts[model], geometries[area], scenario)
        if entry.get("Status") != status or (payload is None) != (entry.get("Maximum Safe Payload (kg)") is None):
            raise ValueError(f"Capacity status mismatch {area}/{model}")
        if payload is not None and not _near(payload, entry["Maximum Safe Payload (kg)"], atol=1e-5):
            raise ValueError(f"Capacity value mismatch {area}/{model}")
        found[(area, model)] = {"Scenario": scenario.key, "rho": scenario.rho,
                                "alpha": scenario.alpha, "beta": scenario.beta, **entry}
    if len(found) != 45:
        raise ValueError("Incomplete capacity model-area pairs")
    return [found[key] for key in sorted(found)]


def _candidate_totals(state: dict, areas: dict) -> dict[str, int]:
    counts = state.get("candidate_counts")
    if not isinstance(counts, dict) or set(counts) != set(areas):
        raise ValueError("Candidate counts missing service areas")
    total = {name: 0 for name in ("possible", "physical", "retained", "dominated_removed")}
    for item in counts.values():
        for name in total:
            value = item.get(name)
            if not isinstance(value, int) or value < 0:
                raise ValueError(f"Invalid candidate count {name}")
            total[name] += value
        if item["physical"] - item["dominated_removed"] != item["retained"] or item["physical"] > item["possible"]:
            raise ValueError("Candidate deletion accounting fails")
    return total


def _validate_epsilon(state: dict, groups, aircrafts, geometries, scenario: Scenario) -> tuple[list[dict], dict[str, int], list[str]]:
    plan = state.get("epsilon_plan", [])
    results = state.get("epsilon_results", {})
    if not isinstance(plan, list) or not isinstance(results, dict):
        raise ValueError("Invalid epsilon plan/results")
    planned = {job["id"]: job for job in plan}
    if len(planned) != len(plan) or set(results) - set(planned):
        raise ValueError("Duplicate/unplanned epsilon job")
    base = [job for job in plan if job.get("stage") == "base"]
    if plan and (len(base) != 363 or any(sum(job["primary"] == k for job in base) != 121 for k in OBJECTIVES)):
        raise ValueError("Initial epsilon grid is not 3 x 11 x 11")
    counts = {"planned": len(plan), "optimal": 0, "infeasible": 0, "tentative": 0}
    rows = []
    errors = []
    positions = {job["id"]: index for index, job in enumerate(plan)}

    def wider(source_limits: dict, current: dict) -> bool:
        return set(source_limits).issubset(current) and all(source_limits[k] >= current[k] for k in source_limits)

    def same_number(a, b) -> bool:
        return (a is None and b is None) or _near(a, b, atol=1e-9, rtol=1e-9)

    for job in plan:
        key, primary = job["id"], job["primary"]
        bounds = job["limits"]
        record = results.get(key)
        status = record.get("status", "missing") if isinstance(record, dict) else "missing"
        if isinstance(record, dict) and status in ("optimal", "infeasible"):
            if record.get("primary") != primary or record.get("limits") != bounds:
                errors.append(f"{key}: objective/limits changed")
                status = "unverified"
            elif record.get("reused") is True:
                proof = record.get("proof_source")
                if proof == f"endpoint:{primary}":
                    source = state.get("endpoints", {}).get(primary)
                    source_limits = {}
                    if source and source.get("status") == "optimal":
                        source = {"status": "optimal", "selected": source["selected"],
                                  "metrics": source["metrics"], "dual_bound": source["primary_dual_bound"],
                                  "gap": 0., "root_proof_source": proof}
                else:
                    source = results.get(proof)
                    source_limits = planned[proof]["limits"] if proof in planned else None
                if (source is None or source_limits is None or
                    (proof in positions and positions[proof] >= positions[key]) or
                    not wider(source_limits, bounds) or source.get("status") != status or
                    record.get("root_proof_source") != source.get("root_proof_source", proof) or
                    record.get("certificate_limits") != source_limits or
                    not same_number(record.get("dual_bound"), source.get("dual_bound")) or
                    not same_number(record.get("gap"), source.get("gap"))):
                    errors.append(f"{key}: invalid certificate containment/provenance")
                    status = "unverified"
                elif status == "optimal" and (
                    sorted(record.get("selected", [])) != sorted(source.get("selected", [])) or
                    record.get("metrics") != source.get("metrics") or
                    any(source["metrics"][j] > limit for j, limit in bounds.items())):
                    errors.append(f"{key}: certified incumbent violates tighter limits")
                    status = "unverified"
            elif record.get("reused") is not False or record.get("proof_source") != key:
                errors.append(f"{key}: direct MILP provenance missing")
                status = "unverified"
        if status == "optimal":
            if record.get("solver_status") != 0 or not _near(record.get("gap"), 0., atol=1e-7):
                errors.append(f"{key}: purported optimal without solver proof")
                status = "unverified"
            else:
                actual, _, _ = _solution(record["selected"], groups, aircrafts, geometries, scenario)
                _assert_metrics(actual, record["metrics"])
                if record.get("primary") != primary or any(actual[j] > v + 1e-6 * max(1., abs(v)) for j, v in bounds.items()):
                    errors.append(f"{key}: objective or epsilon bound mismatch")
                    status = "unverified"
                elif record.get("dual_bound") is None or record["dual_bound"] > actual[primary] + 1e-5 * max(1., actual[primary]):
                    errors.append(f"{key}: invalid optimality bound")
                    status = "unverified"
        elif status == "infeasible" and record.get("solver_status") != 2:
            errors.append(f"{key}: infeasible label not proven")
            status = "unverified"
        if status in ("optimal", "infeasible"):
            counts[status] += 1
        else:
            counts["tentative"] += 1
        values = record.get("metrics") if isinstance(record, dict) and isinstance(record.get("metrics"), dict) else {}
        rows.append({"Scenario": scenario.key, "Job ID": key, "Stage": job["stage"], "Primary": primary,
                     "Bound N": bounds.get("N"), "Bound E (kWh)": bounds.get("E"), "Bound T (s)": bounds.get("T"),
                     "Status": status, "Solver status": record.get("solver_status") if isinstance(record, dict) else None,
                     "Reused certificate": record.get("reused") if isinstance(record, dict) else None,
                     "Proof source": record.get("proof_source") if isinstance(record, dict) else None,
                     "Root proof source": record.get("root_proof_source") if isinstance(record, dict) else None,
                     "Sorties": values.get("N"), "Energy (kWh)": values.get("E"), "Operation Time (s)": values.get("T"),
                     "Primary lower bound": record.get("dual_bound") if isinstance(record, dict) else None,
                     "MIP gap": record.get("gap") if isinstance(record, dict) else None,
                     "Nodes": record.get("node_count") if isinstance(record, dict) else None,
                     "Selected batches": len(record.get("selected", [])) if isinstance(record, dict) else None})
    return rows, counts, errors


def _xlsx_from_template(baseline_rows: list[dict]) -> None:
    """Preserve all six source sheets and write only Q1 values in a new workbook."""
    wb = load_workbook(TEMPLATE)
    try:
        if Q1_SHEET not in wb.sheetnames:
            raise ValueError("Q1 template sheet missing")
        sheet = wb[Q1_SHEET]
        headers = tuple(sheet.cell(1, c).value for c in range(1, 10))
        if headers != Q1_HEADERS:
            raise ValueError(f"Q1 template columns changed: {headers}")
        if any(sheet.cell(r, c).value is not None for r in range(2, sheet.max_row + 1) for c in range(1, 10)):
            raise ValueError("Q1 template already contains result values")
        names = tuple(wb.sheetnames)
        for row_index, data in enumerate(baseline_rows, 2):
            for col_index, header in enumerate(Q1_HEADERS, 1):
                cell = sheet.cell(row_index, col_index)
                if row_index > 3:
                    source = sheet.cell(2, col_index)
                    if source.has_style:
                        cell._style = copy(source._style)
                    if source.number_format:
                        cell.number_format = source.number_format
                cell.value = data[header]
        XLSX_OUT.parent.mkdir(parents=True, exist_ok=True)
        temp = XLSX_OUT.with_name(XLSX_OUT.stem + f".{os.getpid()}.tmp.xlsx")
        try:
            wb.save(temp)
            verification = load_workbook(temp, read_only=True, data_only=False)
            try:
                if tuple(verification.sheetnames) != names:
                    raise ValueError("Workbook sheet topology changed")
                output_sheet = verification[Q1_SHEET]
                if tuple(output_sheet.cell(1, c).value for c in range(1, 10)) != Q1_HEADERS:
                    raise ValueError("Written Q1 header changed")
                if output_sheet.cell(len(baseline_rows) + 1, 1).value != baseline_rows[-1]["Sortie ID"]:
                    raise ValueError("Written Q1 final row missing")
            finally:
                verification.close()
            os.replace(temp, XLSX_OUT)
        finally:
            if temp.exists():
                temp.unlink()
    finally:
        wb.close()


def collect(checkpoint_dir: Path = CHECKPOINT_DIR) -> tuple[dict[str, list[dict]], dict]:
    origin, areas, groups, aircrafts, inputs = load_full_inputs(ROOT)
    fingerprint = make_fingerprint(inputs)
    geometries = None
    tables = {name: [] for name in CSV_COLUMNS}
    for area in sorted(areas):
        node = areas[area]
        tables["q1_full_input_areas.csv"].append({"Service Area ID": area,
            "Longitude (deg)": node.longitude_deg, "Latitude (deg)": node.latitude_deg,
            "Elevation (m)": node.elevation_m, "Cargo count": len(groups[area])})
        for cargo in groups[area]:
            tables["q1_full_input_boxes.csv"].append({"Cargo Box ID": cargo.cargo_id,
                "Service Area ID": area, "Mass (kg)": cargo.mass_kg, "Volume (m³)": cargo.volume_m3})
    states = {}
    inspected = {}
    for scenario in all_scenarios():
        path = checkpoint_dir / f"{scenario.key}.json"
        if not path.exists():
            inspected[scenario.key] = (None, "missing", "checkpoint absent")
            continue
        try:
            state = read_checkpoint(path, fingerprint)
            if state.get("schema") != 2 or state.get("scenario") != {"rho": scenario.rho, "alpha": scenario.alpha, "beta": scenario.beta}:
                raise ValueError("Scenario/schema mismatch")
            if geometries is None:
                geometries = make_geometries(origin, areas, DEM)
            capacities = _capacity_rows(state, scenario, areas, aircrafts, geometries)
            totals = _candidate_totals(state, areas)
            if totals["possible"] != sum(3 * ((1 << len(boxes)) - 1) for boxes in groups.values()):
                raise ValueError("Incomplete local subset enumeration")
            tables["q1_full_capacity_all.csv"].extend(capacities)
            states[scenario.key] = state
            inspected[scenario.key] = (totals, "validated_checkpoint", "")
        except (ValueError, KeyError, TypeError, AssertionError, OSError) as exc:
            inspected[scenario.key] = (None, "invalid", str(exc))

    formal_rows_by_scenario = {}
    representatives = {}
    tentative_representatives = {}
    for scenario in all_scenarios():
        key = scenario.key
        totals, verdict, reason = inspected[key]
        state = states.get(key) if verdict == "validated_checkpoint" else None
        epsilon_counts = {"planned": 0, "optimal": 0, "infeasible": 0, "tentative": 0}
        formal = False
        report_status = verdict
        if state is not None:
            try:
                if geometries is None:
                    raise AssertionError("Missing geometry")
                endpoint_ok = True
                for primary in OBJECTIVES:
                    endpoint = state.get("endpoints", {}).get(primary)
                    if not isinstance(endpoint, dict):
                        endpoint_ok = False
                        continue
                    status = endpoint.get("status", "unknown")
                    if status == "optimal":
                        actual, _, _ = _solution(endpoint["selected"], groups, aircrafts, geometries, scenario)
                        _assert_metrics(actual, endpoint["metrics"])
                        bound = endpoint.get("primary_dual_bound")
                        if bound is None or not _near(bound, actual[primary], atol=1e-4, rtol=1e-6):
                            raise ValueError(f"Endpoint bound mismatch: {primary}")
                        if len(endpoint.get("area_status", {})) != 15 or any(
                                item.get("status") != "optimal" for item in endpoint["area_status"].values()):
                            raise ValueError(f"Endpoint area proof incomplete: {primary}")
                    else:
                        endpoint_ok = False
                    metric = endpoint.get("metrics") if isinstance(endpoint.get("metrics"), dict) else {}
                    tables["q1_full_endpoints.csv"].append({"Scenario": key, "Primary": primary, "Status": status,
                        "Sorties": metric.get("N"), "Energy (kWh)": metric.get("E"),
                        "Operation Time (s)": metric.get("T"), "Primary lower bound": endpoint.get("primary_dual_bound"),
                        "Area count": len(endpoint.get("area_status", {}))})
                epsilon_rows, epsilon_counts, epsilon_errors = _validate_epsilon(state, groups, aircrafts, geometries, scenario)
                tables["q1_full_epsilon_status.csv"].extend(epsilon_rows)
                if epsilon_errors:
                    raise ValueError("; ".join(epsilon_errors[:3]))
                plans = state.get("epsilon_plan", [])
                refinement_done = state.get("epsilon_stage") == "refine"
                if state.get("status") == "complete_sampled" and not refinement_done:
                    raise ValueError("Completed label conflicts with unfinished epsilon refinement")
                for index, point in enumerate(state.get("pareto_sample", []), 1):
                    actual, _, _ = _solution(point["selected"], groups, aircrafts, geometries, scenario)
                    _assert_metrics(actual, point["metrics"])
                    tables["q1_full_pareto_sample.csv"].append({"Scenario": key, "Sample ID": index,
                        "Sample status": "proven scalar; grid completion pending" if state.get("status") != "complete_sampled" else "proven sampled",
                        "Sorties": actual["N"], "Energy (kWh)": actual["E"],
                        "Operation Time (s)": actual["T"], "Selected batch IDs": ";".join(sorted(point["selected"]))})
                formal = (state.get("status") == "complete_sampled" and refinement_done and endpoint_ok and
                          epsilon_counts["planned"] >= 363 and epsilon_counts["tentative"] == 0 and
                          state.get("representative") is not None and state.get("tentative_representative") is None)
                if formal:
                    ids = state["representative"]
                    expected_ids, ideals, upper = _expected_representative(
                        state.get("pareto_sample", []), state["endpoints"])
                    if set(ids) != set(expected_ids):
                        raise ValueError("Representative violates ideal-distance selection")
                    if any(not _near(state.get("ideal", {}).get(j), ideals[j]) or
                           not _near(state.get("sample_upper", {}).get(j), upper[j]) for j in OBJECTIVES):
                        raise ValueError("Recorded normalization interval changed")
                    metric, rows, pairs = _solution(ids, groups, aircrafts, geometries, scenario)
                    representatives[key] = (metric, pairs)
                    formal_rows_by_scenario[key] = rows
                    tables["q1_full_formal_batches.csv"].extend(rows)
                    if not any(set(point["selected"]) == set(ids) for point in state.get("pareto_sample", [])):
                        raise ValueError("Formal representative absent from sampled Pareto set")
                    report_status = "formal_sampled"
                    for row in tables["q1_full_pareto_sample.csv"]:
                        if row["Scenario"] == key:
                            row["Sample status"] = "formal sampled"
                else:
                    report_status = "incomplete_or_tentative"
                    reason = state.get("status", "checkpoint unfinished")
                    provisional = state.get("tentative_representative")
                    if provisional:
                        provisional_metric, _, provisional_pairs = _solution(provisional, groups, aircrafts, geometries, scenario)
                        tentative_representatives[key] = (provisional_metric, provisional_pairs)
            except (ValueError, KeyError, TypeError, AssertionError) as exc:
                report_status, reason, formal = "invalid", str(exc), False
                representatives.pop(key, None)
                tentative_representatives.pop(key, None)
                formal_rows_by_scenario.pop(key, None)
                tables["q1_full_formal_batches.csv"] = [r for r in tables["q1_full_formal_batches.csv"] if r["Scenario"] != key]
                tables["q1_full_pareto_sample.csv"] = [r for r in tables["q1_full_pareto_sample.csv"] if r["Scenario"] != key]
        tables["q1_full_scenario_status.csv"].append({"Scenario": key, "rho": scenario.rho, "alpha": scenario.alpha,
            "beta": scenario.beta, "Checkpoint status": state.get("status") if state else verdict,
            "Report status": report_status, "Formal": formal, "Reason": reason,
            "Candidates possible": totals.get("possible") if totals else None,
            "Candidates physical": totals.get("physical") if totals else None,
            "Candidates retained": totals.get("retained") if totals else None,
            "Dominated removed": totals.get("dominated_removed") if totals else None,
            "Epsilon planned": epsilon_counts["planned"], "Epsilon optimal": epsilon_counts["optimal"],
            "Epsilon infeasible": epsilon_counts["infeasible"], "Epsilon tentative": epsilon_counts["tentative"]})
    baseline = representatives.get(Scenario(.2, 1., 1.).key)
    status_by_key = {r["Scenario"]: r["Report status"] for r in tables["q1_full_scenario_status.csv"]}
    for scenario in all_scenarios():
        key = scenario.key
        current = representatives.get(key)
        if current is None:
            current = tentative_representatives.get(key)
            if current is None:
                tables["q1_full_sensitivity.csv"].append({"Scenario": key, "Report status": status_by_key[key],
                    "Representative type": "none", "Baseline comparison available": baseline is not None})
                continue
        metric, pairs = current
        base_metric, base_pairs = baseline if baseline is not None else (None, None)
        tables["q1_full_sensitivity.csv"].append({"Scenario": key, "Report status": status_by_key[key],
            "Representative type": "formal sampled" if key in representatives else "tentative only",
            "Sorties": metric["N"], "Energy (kWh)": metric["E"],
            "Operation Time (s)": metric["T"],
            "Delta sorties": metric["N"] - base_metric["N"] if base_metric else None,
            "Delta energy (kWh)": metric["E"] - base_metric["E"] if base_metric else None,
            "Delta operation time (s)": metric["T"] - base_metric["T"] if base_metric else None,
            "Changed batch pairs": len(pairs.symmetric_difference(base_pairs)) if base_pairs is not None else None,
            "Baseline comparison available": baseline is not None})
    summary = {"status": "all_27_formal_sampled" if len(representatives) == 27 else "incomplete",
               "scenario_count": 27, "formal_scenarios": len(representatives),
               "missing_or_unproven": [r["Scenario"] for r in tables["q1_full_scenario_status.csv"] if not r["Formal"]],
               "baseline_xlsx_ready": Scenario(.2, 1., 1.).key in formal_rows_by_scenario,
               "existing_xlsx_requires_review": XLSX_OUT.exists() and Scenario(.2, 1., 1.).key not in formal_rows_by_scenario,
               "input_fingerprint": fingerprint, "template_sha256": file_hash(TEMPLATE),
               "generated_utc": datetime.now(timezone.utc).isoformat()}
    return tables, {"summary": summary, "baseline_rows": formal_rows_by_scenario.get(Scenario(.2, 1., 1.).key)}


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate all 27 full Q1 checkpoints and publish formal/partial output")
    parser.add_argument("--check-only", action="store_true", help="inspect checkpoints without writing output")
    args = parser.parse_args()
    tables, detail = collect()
    summary = detail["summary"]
    if not args.check_only:
        for name, columns in CSV_COLUMNS.items():
            write_csv_atomic(TABLE_DIR / name, columns, tables[name])
        if detail["baseline_rows"]:
            _xlsx_from_template(detail["baseline_rows"])
        atomic_json(STATUS_OUT, summary)
    print(json.dumps({"status": summary["status"], "formal_scenarios": summary["formal_scenarios"],
                      "missing_or_unproven": summary["missing_or_unproven"],
                      "baseline_xlsx_ready": summary["baseline_xlsx_ready"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
