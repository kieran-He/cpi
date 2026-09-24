"""Read only the verified Q1 inputs needed by the P1 vertical slice."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


@dataclass(frozen=True)
class Node:
    node_id: str
    longitude_deg: float
    latitude_deg: float
    elevation_m: float


@dataclass(frozen=True)
class Cargo:
    cargo_id: str
    service_area_id: str
    mass_kg: float
    volume_m3: float


@dataclass(frozen=True)
class Aircraft:
    model_id: str
    empty_mass_kg: float
    payload_capacity_kg: float
    loading_volume_m3: float
    cruise_speed_m_s: float
    unloaded_range_m: float
    full_range_m: float
    usable_energy_kwh: float
    reserve_fraction: float
    prep_time_s: float
    loading_time_per_box_s: float
    base_handover_time_s: float
    handover_time_per_box_s: float
    ascent_speed_m_s: float
    descent_speed_m_s: float
    ascent_efficiency: float


def _sheet_rows(path: Path, sheet_name: str) -> tuple[list[str], list[tuple[Any, ...]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"{path.name}: missing sheet {sheet_name!r}")
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    if not header or not any(value is not None for value in header):
        raise ValueError(f"{path.name}/{sheet_name}: blank header")
    names = [str(value).strip() if value is not None else "" for value in header]
    records = [tuple(row) for row in rows if any(value is not None for value in row)]
    return names, records


def _column_map(headers: list[str], required: tuple[str, ...], source: str) -> dict[str, int]:
    mapping = {name: i for i, name in enumerate(headers)}
    missing = [name for name in required if name not in mapping]
    if missing:
        raise ValueError(f"{source}: required columns not found: {missing}")
    return mapping


def load_q1_inputs(project_root: Path) -> dict[str, Any]:
    data_dir = project_root / "questions" / "Data" / "Basic Data for Drone-Based Emergency Supply Transport"
    node_path = data_dir / "Dispatch Center and Service Areas.xlsx"
    cargo_path = data_dir / "Supply Demand and Delivery Deadlines.xlsx"
    aircraft_path = data_dir / "Transport Drone Data.xlsx"

    nodes: dict[str, Node] = {}
    wb = load_workbook(node_path, read_only=True, data_only=True)
    ws = wb["Data"]
    current_columns: dict[str, int] | None = None
    for row in ws.iter_rows(values_only=True):
        values = tuple(row)
        if "Dispatch Center ID" in values or "Service Area ID" in values:
            current_columns = {
                "id": values.index("Dispatch Center ID") if "Dispatch Center ID" in values else values.index("Service Area ID"),
                "lon": next((i for i, value in enumerate(values) if isinstance(value, str) and value.startswith("Longitude")), -1),
                "lat": next((i for i, value in enumerate(values) if isinstance(value, str) and value.startswith("Latitude")), -1),
                "elev": values.index("Elevation (m)") if "Elevation (m)" in values else -1,
            }
            if min(current_columns.values()) < 0:
                raise ValueError(f"{node_path.name}: malformed node header row: {values}")
            continue
        if current_columns is None:
            continue
        ident = values[current_columns["id"]] if current_columns["id"] < len(values) else None
        if ident == "O01" or ident == "S001":
            nodes[str(ident)] = Node(str(ident), float(values[current_columns["lon"]]), float(values[current_columns["lat"]]), float(values[current_columns["elev"]]))
    if set(nodes) != {"O01", "S001"}:
        raise ValueError(f"Could not resolve O01 and S001 coordinates from node workbook: {sorted(nodes)}")

    h, rows = _sheet_rows(cargo_path, "Cargo Box List")
    cc = _column_map(h, ("Cargo Box ID", "Service Area ID", "Mass per Box (kg)", "Volume per Box (m³)"), cargo_path.name)
    cargos_all = [Cargo(str(r[cc["Cargo Box ID"]]), str(r[cc["Service Area ID"]]), float(r[cc["Mass per Box (kg)"]]), float(r[cc["Volume per Box (m³)"]])) for r in rows]
    if len(cargos_all) != 80:
        raise ValueError(f"Expected 80 detailed cargo rows, found {len(cargos_all)}")
    cargos = [cargo for cargo in cargos_all if cargo.service_area_id == "S001"][:3]
    expected_ids = ["S001-MED-01", "S001-MED-02", "S001-WAT-01"]
    if [cargo.cargo_id for cargo in cargos] != expected_ids:
        raise ValueError(f"S001 first-three cargo IDs differ from verified P1 slice: {[c.cargo_id for c in cargos]}")

    aircraft_wb = load_workbook(aircraft_path, read_only=True, data_only=True)
    aircraft_ws = aircraft_wb["Data"]
    aircraft_rows = [tuple(row) for row in aircraft_ws.iter_rows(values_only=True)]
    aircraft_header_row = next((i for i, row in enumerate(aircraft_rows) if "Model ID" in row and "Maximum Cargo Mass (kg)" in row), None)
    if aircraft_header_row is None:
        raise ValueError(f"{aircraft_path.name}: aircraft model table header not found")
    h = [str(value).strip() if value is not None else "" for value in aircraft_rows[aircraft_header_row]]
    rows = aircraft_rows[aircraft_header_row + 1:]
    ac = _column_map(h, ("Model ID", "Total Mass Without Payload, Including Battery (kg)", "Maximum Cargo Mass (kg)", "Usable Loading Volume (m³)", "Planned Cruise Speed (m/s)", "Standard Unloaded Range (m)", "Standard Fully Loaded Range (m)", "Usable Battery Energy (kWh)", "Minimum Return Battery Level (%)", "Fixed Preparation Time at the Workstation (s)", "Loading Time per Box (s)", "Base Handover Time at Receiving Point (s)", "Additional Handover Time per Box (s)", "Maximum Ascent Speed (m/s)", "Maximum Descent Speed (m/s)", "Ascent Energy Efficiency"), aircraft_path.name)
    # Only rows with a single-character A/B/C model ID are aircraft-model records.
    aircraft: dict[str, Aircraft] = {}
    for r in rows:
        model = r[ac["Model ID"]]
        if model not in {"A", "B", "C"}:
            continue
        parameter_columns = [ac[name] for name in ("Total Mass Without Payload, Including Battery (kg)", "Maximum Cargo Mass (kg)", "Usable Loading Volume (m³)", "Planned Cruise Speed (m/s)", "Standard Unloaded Range (m)", "Standard Fully Loaded Range (m)", "Usable Battery Energy (kWh)", "Minimum Return Battery Level (%)")]
        if any(index >= len(r) or r[index] is None for index in parameter_columns):
            continue  # Later workbook sections also contain model IDs but not model parameters.
        aircraft[str(model)] = Aircraft(
            str(model), float(r[ac["Total Mass Without Payload, Including Battery (kg)"]]),
            float(r[ac["Maximum Cargo Mass (kg)"]]), float(r[ac["Usable Loading Volume (m³)"]]),
            float(r[ac["Planned Cruise Speed (m/s)"]]), float(r[ac["Standard Unloaded Range (m)"]]),
            float(r[ac["Standard Fully Loaded Range (m)"]]), float(r[ac["Usable Battery Energy (kWh)"]]),
            float(r[ac["Minimum Return Battery Level (%)"]]) / 100.0,
            float(r[ac["Fixed Preparation Time at the Workstation (s)"]]), float(r[ac["Loading Time per Box (s)"]]),
            float(r[ac["Base Handover Time at Receiving Point (s)"]]), float(r[ac["Additional Handover Time per Box (s)"]]),
            float(r[ac["Maximum Ascent Speed (m/s)"]]), float(r[ac["Maximum Descent Speed (m/s)"]]),
            float(r[ac["Ascent Energy Efficiency"]]),
        )
    if set(aircraft) != {"A", "B", "C"}:
        raise ValueError(f"Expected model rows A/B/C, found {sorted(aircraft)}")
    return {"nodes": nodes, "cargos": cargos, "all_cargo_count": len(cargos_all), "aircraft": aircraft,
            "input_paths": [node_path, cargo_path, aircraft_path]}
