"""Full Q1 model. The existing P1 modules are imported, never changed.

Planning assumptions: the Q1 document's energy split is conditional; alpha/beta
are diagnostic multipliers. Only identical cargo subsets may lose dominated
aircraft choices. Global epsilon constraints couple service areas, so epsilon
MILPs must use the full 80-box incidence matrix. A solver time limit is a
reporting limit, never permission to relabel a feasible incumbent as optimal.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from itertools import product
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import csc_matrix, vstack

from src.data.q1_inputs import Aircraft, Cargo, Node, _column_map, _sheet_rows, load_q1_inputs
from src.models.q1_physics import G_M_S2, route_geometry

OBJECTIVES = ("N", "E", "T")
TOL = 1e-8


@dataclass(frozen=True)
class Scenario:
    rho: float
    alpha: float
    beta: float

    @property
    def key(self) -> str:
        return f"r{self.rho:.2f}_a{self.alpha:.1f}_b{self.beta:.1f}"


@dataclass(frozen=True)
class Geometry:
    distance_m: float
    terrain_max_m: float
    sampled_cells: int
    outbound_climb_m: float
    outbound_descent_m: float
    inbound_climb_m: float
    inbound_descent_m: float


@dataclass(frozen=True)
class Batch:
    batch_id: str
    area_id: str
    model_id: str
    mask: int
    cargo_ids: tuple[str, ...]
    mass_kg: float
    volume_m3: float
    energy_kwh: float
    operation_time_s: float
    round_trip_time_s: float
    return_soc: float


def all_scenarios() -> list[Scenario]:
    return [Scenario(r, a, b) for r, a, b in product((.2, .3, .4), (.9, 1., 1.1), (.9, 1., 1.1))]


def load_full_inputs(root: Path) -> tuple[Node, dict[str, Node], dict[str, list[Cargo]], dict[str, Aircraft], list[Path]]:
    """Reuse verified P1 workbook column/parser logic while reading all rows."""
    p1 = load_q1_inputs(root)  # Also verifies the original three-box P1 slice.
    node_path, cargo_path, aircraft_path = p1["input_paths"]
    wb = load_workbook(node_path, read_only=True, data_only=True)
    nodes: dict[str, Node] = {}
    area_ids = [f"S{i:03d}" for i in range(1, 16)]
    columns: dict[str, int] | None = None
    try:
        for row in wb["Data"].iter_rows(values_only=True):
            vals = tuple(row)
            if "Dispatch Center ID" in vals or "Service Area ID" in vals:
                columns = {
                    "id": vals.index("Dispatch Center ID") if "Dispatch Center ID" in vals else vals.index("Service Area ID"),
                    "lon": next((i for i, v in enumerate(vals) if isinstance(v, str) and v.startswith("Longitude")), -1),
                    "lat": next((i for i, v in enumerate(vals) if isinstance(v, str) and v.startswith("Latitude")), -1),
                    "elev": vals.index("Elevation (m)") if "Elevation (m)" in vals else -1,
                }
                if min(columns.values()) < 0:
                    raise ValueError("Malformed node heading")
                continue
            if columns is None or columns["id"] >= len(vals):
                continue
            ident = vals[columns["id"]]
            if isinstance(ident, str) and (ident == "O01" or ident in area_ids):
                if ident in nodes:
                    raise ValueError(f"Duplicate node {ident}")
                nodes[ident] = Node(ident, float(vals[columns["lon"]]), float(vals[columns["lat"]]), float(vals[columns["elev"]]))
    finally:
        wb.close()
    if set(nodes) != {"O01", *area_ids}:
        raise ValueError(f"Expected O01 and S001-S015; got {sorted(nodes)}")

    headings, rows = _sheet_rows(cargo_path, "Cargo Box List")
    cols = _column_map(headings, ("Cargo Box ID", "Service Area ID", "Mass per Box (kg)", "Volume per Box (m³)"), cargo_path.name)
    groups: dict[str, list[Cargo]] = {area: [] for area in area_ids}
    seen: set[str] = set()
    for row in rows:
        ident, area = str(row[cols["Cargo Box ID"]]), str(row[cols["Service Area ID"]])
        if ident in seen or area not in groups:
            raise ValueError(f"Duplicate cargo or unknown area: {ident}/{area}")
        seen.add(ident)
        mass, volume = float(row[cols["Mass per Box (kg)"]]), float(row[cols["Volume per Box (m³)"]])
        if not (math.isfinite(mass) and math.isfinite(volume) and mass > 0 and volume > 0):
            raise ValueError(f"Invalid cargo quantity: {ident}")
        groups[area].append(Cargo(ident, area, mass, volume))
    if len(seen) != 80 or any(not v for v in groups.values()):
        raise ValueError(f"Expected 80 boxes across 15 nonempty areas; got {len(seen)}")
    for boxes in groups.values():
        boxes.sort(key=lambda c: c.cargo_id)
    return nodes["O01"], {k: nodes[k] for k in area_ids}, groups, p1["aircraft"], [node_path, cargo_path, aircraft_path]


def make_geometries(origin: Node, areas: dict[str, Node], dem_path: Path) -> dict[str, Geometry]:
    geometries = {}
    for area, dst in areas.items():
        out_d, out_z, out_n = route_geometry(origin, dst, str(dem_path))
        in_d, in_z, in_n = route_geometry(dst, origin, str(dem_path))
        if abs(out_d - in_d) > 1e-5 or abs(out_z - in_z) > 1e-5:
            raise ValueError(f"Route geometry differs by direction: {area}")
        cruise = out_z + 50.0
        dst_alt = dst.elevation_m + 30.0
        geometries[area] = Geometry(out_d, out_z, max(out_n, in_n),
                                    max(0., cruise - origin.elevation_m), max(0., cruise - dst_alt),
                                    max(0., cruise - dst_alt), max(0., cruise - origin.elevation_m))
    return geometries


def range_m(aircraft: Aircraft, payload: float) -> float:
    return aircraft.unloaded_range_m - (aircraft.unloaded_range_m - aircraft.full_range_m) * (payload / aircraft.payload_capacity_kg) ** 1.5


def flight_energy(aircraft: Aircraft, geom: Geometry, payload: float, scenario: Scenario) -> tuple[float, float, float]:
    """Return energy, flight time, range-fraction for loaded outbound/empty inbound."""
    if not 0 <= payload <= aircraft.payload_capacity_kg + TOL:
        raise ValueError("Payload outside aircraft capacity")
    d = geom.distance_m
    horizontal = aircraft.usable_energy_kwh * d * (1. / range_m(aircraft, payload) + 1. / aircraft.unloaded_range_m)
    ascent = G_M_S2 / (3.6e6 * aircraft.ascent_efficiency) * (
        (aircraft.empty_mass_kg + payload) * geom.outbound_climb_m + aircraft.empty_mass_kg * geom.inbound_climb_m)
    time = ((geom.outbound_climb_m + geom.inbound_climb_m) / aircraft.ascent_speed_m_s
            + 2. * d / aircraft.cruise_speed_m_s
            + (geom.outbound_descent_m + geom.inbound_descent_m) / aircraft.descent_speed_m_s)
    fraction = d / range_m(aircraft, payload) + d / aircraft.unloaded_range_m
    return scenario.alpha * horizontal + scenario.beta * ascent, time, fraction


def max_safe_payload(aircraft: Aircraft, geom: Geometry, scenario: Scenario) -> tuple[str, float | None]:
    """Energy-only continuous boundary per Q1; volume remains a batch constraint."""
    limit = (1. - scenario.rho) * aircraft.usable_energy_kwh
    def feasible(q: float) -> bool:
        energy, _, _ = flight_energy(aircraft, geom, q, scenario)
        return energy <= limit + 1e-10
    if not feasible(0.):
        return "unreachable_empty", None
    if feasible(aircraft.payload_capacity_kg):
        return "capacity_bound", aircraft.payload_capacity_kg
    lo, hi = 0., aircraft.payload_capacity_kg
    for _ in range(54):
        mid = (lo + hi) / 2.
        if feasible(mid):
            lo = mid
        else:
            hi = mid
    if not feasible(lo):
        raise AssertionError("Payload root failed feasibility recheck")
    return "energy_bound", lo


def enumerate_area(area: str, cargoes: list[Cargo], aircrafts: dict[str, Aircraft], geom: Geometry,
                   scenario: Scenario) -> tuple[list[Batch], dict[str, int]]:
    """Enumerate all nonempty local subsets, then Pareto-delete only within a subset.

    Q2 must regenerate its own unpruned archive if aircraft availability matters.
    """
    n = len(cargoes)
    if n > 15:
        raise ValueError(f"Unexpected >15 cargoes in {area}")
    masses = [0.] * (1 << n)
    volumes = [0.] * (1 << n)
    counts = [0] * (1 << n)
    for mask in range(1, 1 << n):
        bit = mask & -mask
        i = bit.bit_length() - 1
        prev = mask ^ bit
        masses[mask] = masses[prev] + cargoes[i].mass_kg
        volumes[mask] = volumes[prev] + cargoes[i].volume_m3
        counts[mask] = counts[prev] + 1
    physical = 0
    retained: list[Batch] = []
    energy_cache: dict[tuple[str, float], tuple[float, float, float]] = {}
    for mask in range(1, 1 << n):
        subset: list[Batch] = []
        mass, volume, count = masses[mask], volumes[mask], counts[mask]
        ids = tuple(cargoes[i].cargo_id for i in range(n) if mask & (1 << i))
        for model, aircraft in sorted(aircrafts.items()):
            if mass > aircraft.payload_capacity_kg + TOL or volume > aircraft.loading_volume_m3 + TOL:
                continue
            cache_key = (model, mass)
            if cache_key not in energy_cache:
                energy_cache[cache_key] = flight_energy(aircraft, geom, mass, scenario)
            energy, flight_time, range_fraction = energy_cache[cache_key]
            soc = 1. - energy / aircraft.usable_energy_kwh
            if soc + 1e-10 < scenario.rho or range_fraction > 1. + 1e-10:
                continue
            hand = aircraft.base_handover_time_s + count * aircraft.handover_time_per_box_s
            round_trip = flight_time + hand
            operation = round_trip + aircraft.prep_time_s + count * aircraft.loading_time_per_box_s
            subset.append(Batch(f"{area}|{model}|{mask:05d}", area, model, mask, ids, mass, volume,
                                energy, operation, round_trip, soc))
        physical += len(subset)
        for candidate in subset:
            if not any(other.batch_id != candidate.batch_id and
                       other.energy_kwh <= candidate.energy_kwh + 1e-10 and
                       other.operation_time_s <= candidate.operation_time_s + 1e-8 and
                       (other.energy_kwh < candidate.energy_kwh - 1e-10 or
                        other.operation_time_s < candidate.operation_time_s - 1e-8 or
                        other.batch_id < candidate.batch_id)
                       for other in subset):
                retained.append(candidate)
    return retained, {"possible": len(aircrafts) * ((1 << n) - 1), "physical": physical,
                      "retained": len(retained), "dominated_removed": physical - len(retained)}


def incidence(cargoes: list[Cargo], batches: list[Batch]) -> csc_matrix:
    idx = {c.cargo_id: i for i, c in enumerate(cargoes)}
    rr, cc = [], []
    for j, batch in enumerate(batches):
        for cargo_id in batch.cargo_ids:
            rr.append(idx[cargo_id]); cc.append(j)
    return csc_matrix((np.ones(len(rr)), (rr, cc)), shape=(len(cargoes), len(batches)))


def objectives(batches: list[Batch]) -> dict[str, np.ndarray]:
    return {"N": np.ones(len(batches)), "E": np.array([b.energy_kwh for b in batches]),
            "T": np.array([b.operation_time_s for b in batches])}


def metrics(selected: list[Batch]) -> dict[str, float]:
    return {"N": len(selected), "E": math.fsum(b.energy_kwh for b in selected),
            "T": math.fsum(b.operation_time_s for b in selected)}


def epsilon_reduce_candidates(cargoes: list[Cargo], batches: list[Batch],
                              limits: dict[str, float]) -> tuple[list[Batch], dict[str, int]]:
    """Safely fix candidates that cannot fit an epsilon budget.

    Every selected batch's cost is at least the sum of per-cargo shares
    ``min(cost(batch) / size(batch))`` for the cargoes it carries. The sum of
    these shares is therefore an admissible lower bound for every exact cover.
    Fixing one candidate lets us add the lower bounds of all uncovered cargoes
    and all other service areas; a candidate exceeding any global cap cannot
    occur in a feasible solution and may be removed without changing the MILP.
    """
    if not limits or not batches:
        return batches, {"before": len(batches), "after": len(batches), "removed": 0}

    area_cargoes: dict[str, list[Cargo]] = {}
    for cargo in cargoes:
        area_cargoes.setdefault(cargo.service_area_id, []).append(cargo)
    area_batches: dict[str, list[Batch]] = {area: [] for area in area_cargoes}
    for batch in batches:
        if batch.area_id not in area_batches:
            raise ValueError(f"Batch references unknown area: {batch.area_id}")
        area_batches[batch.area_id].append(batch)

    # A valid lower bound on the cost of covering each area's cargoes.
    cargo_share: dict[str, dict[str, list[float]]] = {}
    area_lower: dict[str, dict[str, float]] = {}
    for area, local_cargoes in area_cargoes.items():
        local_cargoes.sort(key=lambda c: c.cargo_id)
        minima = {name: [math.inf] * len(local_cargoes) for name in limits}
        for batch in area_batches[area]:
            size = batch.mask.bit_count()
            if size == 0:
                raise ValueError(f"Empty batch is invalid: {batch.batch_id}")
            for i in range(len(local_cargoes)):
                if batch.mask & (1 << i):
                    for name in limits:
                        cost = (1. if name == "N" else
                                batch.energy_kwh if name == "E" else
                                batch.operation_time_s if name == "T" else None)
                        if cost is None:
                            raise ValueError(f"Unknown epsilon objective: {name}")
                        minima[name][i] = min(minima[name][i], cost / size)
        if any(not math.isfinite(value) for values in minima.values() for value in values):
            raise ValueError(f"A cargo has no candidate in {area}")
        cargo_share[area] = minima
        area_lower[area] = {name: math.fsum(values) for name, values in minima.items()}

    other_area_lower = {
        area: {name: math.fsum(area_lower[other][name] for other in area_cargoes if other != area)
               for name in limits}
        for area in area_cargoes
    }
    kept: list[Batch] = []
    for batch in batches:
        local_cargoes = area_cargoes[batch.area_id]
        for name, bound in limits.items():
            cost = (1. if name == "N" else
                    batch.energy_kwh if name == "E" else
                    batch.operation_time_s if name == "T" else None)
            if cost is None:
                raise ValueError(f"Unknown epsilon objective: {name}")
            uncovered_lower = math.fsum(
                cargo_share[batch.area_id][name][i]
                for i in range(len(local_cargoes)) if not batch.mask & (1 << i))
            lower = cost + uncovered_lower + other_area_lower[batch.area_id][name]
            if lower > bound + 1e-8 * max(1., abs(bound)):
                break
        else:
            kept.append(batch)
    return kept, {"before": len(batches), "after": len(kept), "removed": len(batches) - len(kept)}


def solve_partition(cargoes: list[Cargo], batches: list[Batch], primary: str,
                    limits: dict[str, float] | None, time_limit_s: float | None) -> dict:
    """One exact MILP call; ``None`` or ``0`` means no solver time limit."""
    if primary not in OBJECTIVES or (time_limit_s is not None and time_limit_s < 0):
        raise ValueError("Invalid primary objective or time limit")
    if not batches:
        return {"status": "infeasible", "selected": [], "metrics": None, "dual_bound": None, "gap": None}
    limits = limits or {}
    batches, reduction = epsilon_reduce_candidates(cargoes, batches, limits)
    if not batches:
        return {"status": "infeasible", "selected": [], "metrics": None, "dual_bound": None,
                "gap": None, "candidate_reduction": reduction}
    matrix = incidence(cargoes, batches)
    vectors = objectives(batches)
    constraints = [matrix]
    lower = [np.ones(len(cargoes))]
    upper = [np.ones(len(cargoes))]
    for name, value in limits.items():
        if name not in OBJECTIVES or not math.isfinite(value):
            raise ValueError(f"Invalid epsilon bound: {name}={value}")
        constraints.append(csc_matrix(vectors[name].reshape(1, -1)))
        lower.append(np.array([-np.inf])); upper.append(np.array([value]))
    mat = vstack(constraints, format="csc")
    options = {"disp": False, "mip_rel_gap": 0.0}
    if time_limit_s not in (None, 0):
        options["time_limit"] = time_limit_s
    res = milp(c=vectors[primary], integrality=np.ones(len(batches)),
               bounds=Bounds(np.zeros(len(batches)), np.ones(len(batches))),
               constraints=LinearConstraint(mat, np.concatenate(lower), np.concatenate(upper)),
               options=options)
    selected: list[Batch] = []
    actual = None
    if res.x is not None and np.all(np.isfinite(res.x)) and np.max(np.abs(res.x - np.rint(res.x))) < 1e-6:
        selected = [b for v, b in zip(res.x, batches) if v > .5]
        counts = {c.cargo_id: 0 for c in cargoes}
        for batch in selected:
            for ident in batch.cargo_ids:
                counts[ident] += 1
        actual = metrics(selected)
        if any(v != 1 for v in counts.values()) or any(actual[k] > bound + 1e-6 * max(1., abs(bound)) for k, bound in limits.items()):
            selected = []; actual = None
    status = "optimal" if res.status == 0 and actual is not None else (
        "infeasible" if res.status == 2 else "time_limit_or_unknown")
    if res.status == 0 and actual is None:
        status = "invalid_incumbent"
    return {"status": status, "solver_status": int(res.status), "message": str(res.message),
            "selected": [b.batch_id for b in selected], "metrics": actual,
            "dual_bound": finite_or_none(getattr(res, "mip_dual_bound", None)),
            "gap": finite_or_none(getattr(res, "mip_gap", None)),
            "node_count": finite_or_none(getattr(res, "mip_node_count", None)),
            "primary": primary, "limits": limits, "candidate_reduction": reduction}


def finite_or_none(value: object) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def validate_selected(ids: list[str], batch_map: dict[str, Batch], cargoes: list[Cargo],
                      aircrafts: dict[str, Aircraft], geoms: dict[str, Geometry], scenario: Scenario) -> dict[str, float]:
    """Recalculate physical and coverage checks from selected batch records."""
    counts = {c.cargo_id: 0 for c in cargoes}
    selected = []
    for ident in ids:
        b = batch_map[ident]
        a = aircrafts[b.model_id]
        e, t, fraction = flight_energy(a, geoms[b.area_id], b.mass_kg, scenario)
        if abs(e - b.energy_kwh) > 1e-8 or fraction > 1. + 1e-10 or 1. - e / a.usable_energy_kwh < scenario.rho - 1e-10:
            raise AssertionError(f"Physical validation failed: {ident}")
        if b.mass_kg > a.payload_capacity_kg + TOL or b.volume_m3 > a.loading_volume_m3 + TOL:
            raise AssertionError(f"Capacity validation failed: {ident}")
        if abs(t + a.base_handover_time_s + len(b.cargo_ids) * a.handover_time_per_box_s + a.prep_time_s + len(b.cargo_ids) * a.loading_time_per_box_s - b.operation_time_s) > 1e-7:
            raise AssertionError(f"Time validation failed: {ident}")
        for cid in b.cargo_ids:
            counts[cid] += 1
        selected.append(b)
    if any(v != 1 for v in counts.values()):
        raise AssertionError("Cargo coverage is not exactly once")
    return metrics(selected)


def nondominated(solutions: list[dict]) -> list[dict]:
    """Finite sampled nondominated set; integer N is compared exactly."""
    by_id: dict[tuple[str, ...], dict] = {}
    for solution in solutions:
        if solution.get("status") == "optimal" and solution.get("metrics") is not None:
            key = tuple(sorted(solution["selected"]))
            by_id[key] = solution
    vals = list(by_id.values())
    def le(a: dict, b: dict) -> bool:
        return a["N"] <= b["N"] and a["E"] <= b["E"] + 1e-8 and a["T"] <= b["T"] + 1e-6
    def strict(a: dict, b: dict) -> bool:
        return a["N"] < b["N"] or a["E"] < b["E"] - 1e-8 or a["T"] < b["T"] - 1e-6
    front = [x for x in vals if not any(le(y["metrics"], x["metrics"]) and strict(y["metrics"], x["metrics"]) for y in vals if y is not x)]
    # Equal objective triples represent one sampled Pareto point; keep a stable ID.
    unique: list[dict] = []
    for x in sorted(front, key=lambda item: tuple(sorted(item["selected"]))):
        m = x["metrics"]
        if not any(m["N"] == y["metrics"]["N"] and abs(m["E"] - y["metrics"]["E"]) <= 1e-8 and
                   abs(m["T"] - y["metrics"]["T"]) <= 1e-6 for y in unique):
            unique.append(x)
    return sorted(unique, key=lambda x: (x["metrics"]["N"], x["metrics"]["E"], x["metrics"]["T"], tuple(sorted(x["selected"]))))


def _pareto_keep_indices(points: np.ndarray) -> list[int]:
    """Return deterministic indices of the non-dominated (N, E, T) rows."""
    if points.size == 0:
        return []
    keep: list[int] = []
    prior_energy = np.empty(0, dtype=float)
    prior_time = np.empty(0, dtype=float)
    for sorties in np.unique(points[:, 0]):
        ids = np.flatnonzero(points[:, 0] == sorties)
        order = ids[np.lexsort((points[ids, 2], points[ids, 1]))]
        group_front: list[int] = []
        best_time = math.inf
        for idx in order:
            op_time = float(points[idx, 2])
            if op_time < best_time - 1e-6:
                group_front.append(int(idx))
                best_time = op_time
        if prior_energy.size:
            local = points[group_front]
            positions = np.searchsorted(prior_energy, local[:, 1] + 1e-8, side="right") - 1
            has_prior = positions >= 0
            dominated = np.zeros(len(group_front), dtype=bool)
            dominated[has_prior] = prior_time[positions[has_prior]] <= local[has_prior, 2] + 1e-6
            group_front = [idx for idx, is_dominated in zip(group_front, dominated) if not is_dominated]
        keep.extend(group_front)
        if group_front:
            combined = np.concatenate((
                np.column_stack((prior_energy, prior_time)) if prior_energy.size else np.empty((0, 2)),
                points[group_front, 1:3],
            ))
            by_energy = combined[np.argsort(combined[:, 0], kind="mergesort")]
            prefix_time = np.minimum.accumulate(by_energy[:, 1])
            # Keep the prefix envelope; searchsorted then checks every earlier N.
            prior_energy = by_energy[:, 0]
            prior_time = prefix_time
    if not keep:
        return []
    keep.sort(key=lambda i: (points[i, 0], points[i, 1], points[i, 2], i))
    unique: list[int] = []
    for idx in keep:
        if not unique:
            unique.append(idx)
            continue
        prior = unique[-1]
        if (points[idx, 0] == points[prior, 0] and
                abs(points[idx, 1] - points[prior, 1]) <= 1e-8 and
                abs(points[idx, 2] - points[prior, 2]) <= 1e-6):
            continue
        unique.append(idx)
    return unique


def exact_area_pareto(cargoes: list[Cargo], batches: list[Batch]) -> tuple[np.ndarray, list[tuple[str, ...]], int]:
    """Enumerate an area's exact-cover Pareto front by subset dynamic programming.

    The least uncovered cargo is assigned next, so every set partition is
    generated once. Dominated partial covers of the same cargo mask can be
    discarded because every continuation adds the same objective vector.
    """
    n = len(cargoes)
    if n > 15:
        raise ValueError(f"Unexpected >15 cargoes in {cargoes[0].service_area_id if cargoes else 'area'}")
    size = 1 << n
    by_mask: dict[int, list[Batch]] = {}
    for batch in sorted(batches, key=lambda item: item.batch_id):
        by_mask.setdefault(batch.mask, []).append(batch)
    fronts: list[np.ndarray | None] = [None] * size
    parents: list[list[tuple[int, int, str]] | None] = [None] * size
    fronts[0] = np.zeros((1, 3), dtype=float)
    parents[0] = []
    total_states = 0
    for mask in range(1, size):
        first = mask & -mask
        sub = mask
        blocks: list[np.ndarray] = []
        choices: list[tuple[int, int, str]] = []
        while sub:
            options = by_mask.get(sub) if sub & first else None
            if options:
                rest = mask ^ sub
                base = fronts[rest]
                assert base is not None
                for batch in options:
                    blocks.append(base + np.array([1., batch.energy_kwh, batch.operation_time_s]))
                    choices.extend((rest, i, batch.batch_id) for i in range(len(base)))
            sub = (sub - 1) & mask
        if blocks:
            points = np.concatenate(blocks, axis=0)
            keep = _pareto_keep_indices(points)
            fronts[mask] = points[keep]
            parents[mask] = [choices[i] for i in keep]
            total_states += len(keep)
        else:
            fronts[mask] = np.empty((0, 3), dtype=float)
            parents[mask] = []

    full = size - 1
    result = fronts[full]
    assert result is not None
    paths: list[tuple[str, ...]] = []
    for final_idx in range(len(result)):
        mask, idx = full, final_idx
        selected: list[str] = []
        while mask:
            parent_rows = parents[mask]
            assert parent_rows is not None
            rest, prior_idx, batch_id = parent_rows[idx]
            selected.append(batch_id)
            mask, idx = rest, prior_idx
        paths.append(tuple(sorted(selected)))
    return result, paths, total_states


def exact_global_pareto(groups: dict[str, list[Cargo]], area_batches: dict[str, list[Batch]],
                        progress=None) -> tuple[list[dict], dict[str, dict[str, float]], dict[str, int]]:
    """Combine exact service-area fronts into the exact fleet-wide Pareto front."""
    points = np.zeros((1, 3), dtype=float)
    paths: list[tuple[str, ...]] = [()]
    area_minima: dict[str, dict[str, float]] = {}
    dp_states: dict[str, int] = {}
    for area in sorted(groups):
        local, local_paths, states = exact_area_pareto(groups[area], area_batches[area])
        if not len(local):
            return [], area_minima, dp_states
        area_minima[area] = {name: float(np.min(local[:, i])) for i, name in enumerate(OBJECTIVES)}
        dp_states[area] = states
        combined = (points[:, None, :] + local[None, :, :]).reshape(-1, 3)
        combined_paths = [left + right for left in paths for right in local_paths]
        keep = _pareto_keep_indices(combined)
        points = combined[keep]
        paths = [combined_paths[i] for i in keep]
        if progress is not None:
            progress(area, len(local), len(points), states)
    result = []
    for row, selected in zip(points, paths):
        result.append({"selected": list(selected),
                       "metrics": {"N": int(round(float(row[0]))), "E": float(row[1]), "T": float(row[2])}})
    result.sort(key=lambda item: (item["metrics"]["N"], item["metrics"]["E"],
                                  item["metrics"]["T"], tuple(item["selected"])))
    return result, area_minima, dp_states


def representative(front: list[dict], ideals: dict[str, float]) -> tuple[dict, dict[str, float]]:
    if not front:
        raise ValueError("No proven sampled solutions")
    ceilings = {j: max(x["metrics"][j] for x in front) for j in OBJECTIVES}
    def distance(x: dict) -> float:
        return math.sqrt(sum(((x["metrics"][j] - ideals[j]) / (ceilings[j] - ideals[j])) ** 2
                             for j in OBJECTIVES if ceilings[j] > ideals[j] + TOL))
    chosen = min(front, key=lambda x: (round(distance(x), 9), x["metrics"]["N"],
                                       x["metrics"]["E"], x["metrics"]["T"], tuple(sorted(x["selected"]))))
    return chosen, ceilings
